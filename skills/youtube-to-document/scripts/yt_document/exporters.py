"""Export validated transcripts as TXT, DOCX, or PDF."""

from io import BytesIO
from math import isfinite
from pathlib import Path

from yt_document.cleanup import CleanedDocument, CleanedParagraph


def format_timestamp(seconds: float) -> str:
    """Format nonnegative seconds as HH:MM:SS."""
    if not isfinite(seconds) or seconds < 0:
        raise ValueError("Timestamp must be finite and nonnegative.")

    total_seconds = int(seconds)
    hours, remainder = divmod(total_seconds, 3600)
    minutes, seconds_remaining = divmod(remainder, 60)

    return f"{hours:02d}:{minutes:02d}:{seconds_remaining:02d}"


def paragraph_timestamp(paragraph: CleanedParagraph) -> str:
    """Format the source time range consistently across exporters."""
    start = format_timestamp(paragraph.start)
    end = format_timestamp(paragraph.end)
    return f"[{start} - {end}]"


def _write_output(
    content: bytes,
    destination: Path,
    extension: str,
) -> Path:
    """Write generated content without replacing an existing file."""
    destination = destination.expanduser().resolve()

    if destination.suffix.lower() != extension:
        raise ValueError(f"Output must use the {extension} extension.")

    destination.parent.mkdir(parents=True, exist_ok=True)

    with destination.open("xb") as handle:
        handle.write(content)

    return destination


def render_txt(document: CleanedDocument) -> str:
    """Render document content without writing a file."""
    lines = [
        document.title,
        f"Source: {document.source_url}",
        f"Language: {document.language}",
        "",
    ]

    for paragraph in document.paragraphs:
        if paragraph.heading:
            lines.extend([paragraph.heading, ""])

        lines.extend(
            [
                paragraph_timestamp(paragraph),
                paragraph.text,
                "",
            ]
        )

    return "\n".join(lines)


def export_txt(document: CleanedDocument, destination: Path) -> Path:
    """Save a UTF-8 text document."""
    content = render_txt(document).encode("utf-8")
    return _write_output(content, destination, ".txt")


def export_docx(document: CleanedDocument, destination: Path) -> Path:
    """Save an editable Word document."""
    from docx import Document
    from docx.shared import Inches, Pt

    word_document = Document()

    normal_style = word_document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.space_after = Pt(8)
    normal_style.paragraph_format.line_spacing = 1.15

    section = word_document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    word_document.core_properties.title = document.title

    word_document.add_heading(document.title, level=0)
    word_document.add_paragraph(f"Source: {document.source_url}")
    word_document.add_paragraph(f"Language: {document.language}")

    for paragraph in document.paragraphs:
        if paragraph.heading:
            word_document.add_heading(paragraph.heading, level=1)

        timestamp = word_document.add_paragraph()
        timestamp.paragraph_format.keep_with_next = True
        timestamp.paragraph_format.space_after = Pt(3)

        timestamp_run = timestamp.add_run(paragraph_timestamp(paragraph))
        timestamp_run.font.size = Pt(9)
        timestamp_run.italic = True

        word_document.add_paragraph(paragraph.text)

    with BytesIO() as buffer:
        word_document.save(buffer)
        content = buffer.getvalue()

    return _write_output(content, destination, ".docx")


def export_pdf(
    document: CleanedDocument,
    destination: Path,
    *,
    font_path: Path | None = None,
) -> Path:
    """Save a PDF with an embedded font and validated character coverage."""
    from fontTools.ttLib import TTFont
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    if font_path is None:
        skill_directory = Path(__file__).resolve().parents[2]
        font_path = (
            skill_directory
            / "assets"
            / "fonts"
            / "NotoSans-Regular.ttf"
        )

    font_path = font_path.expanduser().resolve()

    if not font_path.is_file():
        raise FileNotFoundError(f"PDF font not found: {font_path}")

    # Each block contains text, font size, and line height in millimeters.
    blocks: list[tuple[str, int, int]] = [
        (document.title, 20, 10),
        (f"Source: {document.source_url}", 9, 5),
        (f"Language: {document.language}", 9, 5),
    ]

    for paragraph in document.paragraphs:
        if paragraph.heading:
            blocks.append((paragraph.heading, 14, 7))

        blocks.extend(
            [
                (paragraph_timestamp(paragraph), 9, 5),
                (paragraph.text, 11, 6),
            ]
        )

    with TTFont(str(font_path)) as font:
        supported_characters = set(font.getBestCmap() or {})

    required_characters = {
        ord(character)
        for text, _, _ in blocks
        for character in text
        if character not in "\n\r\t"
    }

    missing = sorted(required_characters - supported_characters)
    if missing:
        preview = ", ".join(f"U+{code:04X}" for code in missing[:8])
        raise ValueError(
            f"The PDF font does not support these characters: {preview}. "
            "Use a font that covers the transcript's language."
        )

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_font("Transcript", fname=str(font_path))
    pdf.set_title(document.title)
    pdf.add_page()

    for text, font_size, line_height in blocks:
        pdf.set_font("Transcript", size=font_size)

        normalized_text = (
            text.replace("\r\n", "\n")
            .replace("\r", "\n")
            .replace("\t", "    ")
        )

        pdf.multi_cell(
            w=0,
            h=line_height,
            text=normalized_text,
            new_x=XPos.LMARGIN,
            new_y=YPos.NEXT,
        )
        pdf.ln(2)

    return _write_output(bytes(pdf.output()), destination, ".pdf")