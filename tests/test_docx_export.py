"""Tests for Word document export."""

from pathlib import Path

import pytest
from docx import Document

from yt_document.cleanup import CleanedDocument, CleanedParagraph
from yt_document.exporters import export_docx


@pytest.fixture
def document() -> CleanedDocument:
    return CleanedDocument(
        title="Sample lesson",
        source_url="https://www.youtube.com/watch?v=abcdefghijk",
        language="en",
        paragraphs=(
            CleanedParagraph(
                heading="Results",
                text="We used 15 samples at the café.",
                segment_ids=(0,),
                start=1.5,
                end=65.2,
            ),
        ),
    )


def test_exports_content_and_heading_structure(
    tmp_path: Path,
    document: CleanedDocument,
) -> None:
    saved = export_docx(document, tmp_path / "exports" / "lesson.docx")
    reopened = Document(str(saved))

    assert [paragraph.text for paragraph in reopened.paragraphs] == [
        "Sample lesson",
        "Source: https://www.youtube.com/watch?v=abcdefghijk",
        "Language: en",
        "Results",
        "[00:00:01 - 00:01:05]",
        "We used 15 samples at the café.",
    ]
    assert reopened.paragraphs[0].style.name == "Title"
    assert reopened.paragraphs[3].style.name == "Heading 1"


def test_does_not_overwrite_existing_file(
    tmp_path: Path,
    document: CleanedDocument,
) -> None:
    destination = tmp_path / "lesson.docx"
    destination.write_bytes(b"Keep this file.")

    with pytest.raises(FileExistsError):
        export_docx(document, destination)

    assert destination.read_bytes() == b"Keep this file."


def test_rejects_wrong_extension(
    tmp_path: Path,
    document: CleanedDocument,
) -> None:
    with pytest.raises(ValueError, match="extension"):
        export_docx(document, tmp_path / "lesson.txt")